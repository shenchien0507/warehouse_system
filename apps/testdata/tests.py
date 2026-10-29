import zipfile
from collections import defaultdict
from io import BytesIO
from unittest import mock

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.testing import Level, make_user
from apps.orders.importer import read_rows

from .generator import COLUMNS, MASK

CATALOG = {
    "platforms": [
        {"name": "蝦皮購物", "orders": 90, "shops": ["蝦皮購物_0號店", "蝦皮購物_1號店"],
         "ship_types": [["蝦皮無包裝", 60], ["蝦皮店到店", 30]], "date_only": False, "order_no_shapes": ["######NAN#D#W#"]},
        {"name": "friDay商城", "orders": 10, "shops": ["friDay商城_3號店"],
         "ship_types": [["7-11", 10]], "date_only": True, "order_no_shapes": ["##############"]},
    ],
    "shipment_shapes": {"蝦皮無包裝": ["TW############U"]},
    "products": [["DO528", "手提雙飲保溫杯", "天空藍-F", 187], ["DO601", "TF高速記憶卡", "128G-F", 793],
                 ["DO479", "防誤食安全餌劑盒", "單入-F", 370]],
}


def form_data(**overrides):
    data = {"first_order_no": 1000}
    for key in ("urgent", "shipping", "new", "formal"):
        data.update({f"{key}_count": 20, f"{key}_start_date": "2026-10-05", f"{key}_start_time": "08:00",
                     f"{key}_end_date": "2026-10-05", f"{key}_end_time": "14:00"})
    data.update({"p0": 90, "p0_t0": 60, "p0_t1": 40, "p1": 10, "p1_t0": 100})
    data.update(overrides)
    return data


@mock.patch("apps.testdata.forms.load_catalog", lambda: CATALOG)
@mock.patch("apps.testdata.generator.load_catalog", lambda: CATALOG)
class GenerateTests(TestCase):
    def setUp(self):
        self.client.login(username=make_user("boss", Level.ADMIN).username, password="pw12345!")

    def post(self, **overrides):
        return self.client.post(reverse("testdata:generate"), form_data(**overrides))

    def files(self, response):
        archive = zipfile.ZipFile(BytesIO(response.content))
        result = {}
        for name in archive.namelist():
            sheet = load_workbook(BytesIO(archive.read(name)), read_only=True).worksheets[0]
            rows = list(sheet.iter_rows(values_only=True))
            result[name.split("_")[0]] = (rows[0], [dict(zip(rows[0], r)) for r in rows[1:]])
        return result

    def test_page_is_admin_only(self):
        self.assertEqual(self.client.get(reverse("testdata:generate")).status_code, 200)
        self.client.login(username=make_user("staff", Level.STAFF).username, password="pw12345!")
        self.assertEqual(self.client.get(reverse("testdata:generate")).status_code, 403)

    def test_generates_one_file_per_type_with_history_columns(self):
        response = self.post()
        self.assertEqual(response["Content-Type"], "application/zip")
        files = self.files(response)
        self.assertEqual(set(files), {"急單", "出貨中", "新訂單", "正式訂單"})
        orders = defaultdict(set)
        for label, (headers, rows) in files.items():
            self.assertEqual(list(headers), COLUMNS)
            for row in rows:
                orders[label].add(row["交易序號"])
                for column in ("買家姓名", "收件者姓名", "收件者手機", "收件者地址"):
                    self.assertEqual(row[column], MASK)
                self.assertIsNone(row["出貨日期"])
                self.assertEqual(bool(row["出貨單號"]), label == "正式訂單")
                if row["交易平台"] == "friDay商城":
                    self.assertEqual(row["建立時間"], "2026/10/05 00:00:00")
                    self.assertEqual(row["出貨類型"], "7-11")
                else:
                    self.assertTrue("2026/10/05 08:00:00" <= row["建立時間"] <= "2026/10/05 14:00:59")
                    self.assertIn(row["出貨類型"], ("蝦皮無包裝", "蝦皮店到店"))
                    self.assertIn(row["交易店鋪"], ("蝦皮購物_0號店", "蝦皮購物_1號店"))
                self.assertEqual(row["商品小計"], row["商品單價"] * row["小計數量"])
        self.assertTrue(all(len(numbers) == 20 for numbers in orders.values()))
        all_numbers = set().union(*orders.values())
        self.assertEqual(all_numbers, set(range(1000, 1080)))  # 4 種類型接續編號、不重複

    def test_zero_ratio_and_zero_count_are_left_out(self):
        files = self.files(self.post(p1=0, p0_t1=0, urgent_count=0, new_count=0))
        self.assertEqual(set(files), {"出貨中", "正式訂單"})
        for _, rows in files.values():
            self.assertTrue(all(r["交易平台"] == "蝦皮購物" and r["出貨類型"] == "蝦皮無包裝" for r in rows))

    def test_file_reads_with_order_importer(self):
        archive = zipfile.ZipFile(BytesIO(self.post().content))
        upload = BytesIO(archive.read(archive.namelist()[0]))
        upload.name = "急單.xlsx"
        headers, rows = read_rows(upload, 1)
        self.assertEqual(headers, COLUMNS)
        self.assertTrue(list(rows))

    def test_invalid_settings_show_errors(self):
        response = self.post(shipping_end_time="07:00")
        self.assertContains(response, "出貨中：結束時間要晚於開始時間。")
        response = self.post(p0=0, p1=0)
        self.assertContains(response, "平台比例至少要有一個大於 0。")
        response = self.post(p0_t0=0, p0_t1=0)
        self.assertContains(response, "蝦皮購物：平台比例大於 0，出貨類型的比例至少要有一個大於 0。")
        response = self.post(urgent_count=0, shipping_count=0, new_count=0, formal_count=0)
        self.assertContains(response, "至少要有一種類型的訂單數大於 0。")
