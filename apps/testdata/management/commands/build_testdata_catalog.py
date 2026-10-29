"""從範例訂單檔抽出「測試資料產生」用的選項，寫進 apps/testdata/catalog.json。

只抽平台、店鋪、出貨類型、商品（編號、名稱、樣式、單價）和單號的格式；買家姓名、電話、地址這些個資一律不讀。
各檔案的訂單量差很多，所以平台比例是每個檔案各自算成百分比再平均（每個檔案份量一樣）；
出貨類型比例用「最新」（建立時間最晚）且有這個平台的那個檔案，例如蝦皮的無包裝是後來才有的，不能被舊資料稀釋。
有新的範例檔時重跑一次即可：python manage.py build_testdata_catalog 歷史紀錄.xlsx ALL.xlsx
大檔（例如 ALL.xlsx）要讀好幾分鐘。"""

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from apps.testdata.catalog import CATALOG_PATH

PRODUCT_LIMIT = 500
SHAPE_LIMIT = 3


def shape(value):
    """把單號的數字換成 #，只留下格式，生成時再填亂數。"""
    return re.sub(r"\d", "#", value)


class Command(BaseCommand):
    help = "從範例訂單檔（歷史紀錄的欄位格式）抽出測試資料產生要用的平台、店鋪、出貨類型、商品清單。"

    def add_arguments(self, parser):
        parser.add_argument("files", nargs="+", help="範例 .xlsx 檔，可以放好幾個")

    def handle(self, *args, files, **options):
        file_stats = []  # 每個檔案：(最晚建立時間, {平台: 訂單數}, {平台: Counter(出貨類型)})
        shops = defaultdict(Counter)
        date_only = defaultdict(Counter)  # 平台 → {True/False: 訂單數}，建立時間是不是都只有日期（00:00:00）
        order_no_shapes = defaultdict(Counter)
        shipment_shapes = defaultdict(Counter)
        products = Counter()
        prices = defaultdict(Counter)

        for path in files:
            self.stdout.write(f"讀取 {path} …")
            sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            index = {str(h).strip(): i for i, h in enumerate(next(rows)) if h is not None}

            def get(row, column):
                i = index.get(column)
                value = row[i] if i is not None and i < len(row) else None
                return "" if value is None else str(value).strip()

            seen = set()
            orders, ship_types, latest = Counter(), defaultdict(Counter), ""
            for row in rows:
                platform = get(row, "交易平台")
                if not platform:
                    continue
                product = (get(row, "商品編號"), get(row, "商品名稱"), get(row, "商品樣式"))
                if product[0]:
                    products[product] += 1
                    price = get(row, "商品單價")
                    if price not in ("", "0"):
                        prices[product][price] += 1
                order_no = get(row, "交易序號")
                if (path, order_no) in seen:
                    continue
                seen.add((path, order_no))
                orders[platform] += 1
                shops[platform][get(row, "交易店鋪")] += 1
                ship_type = get(row, "出貨類型")
                ship_types[platform][ship_type] += 1
                latest = max(latest, get(row, "建立時間"))
                date_only[platform][get(row, "建立時間").endswith("00:00:00")] += 1
                if get(row, "平台訂單編號"):
                    order_no_shapes[platform][shape(get(row, "平台訂單編號"))] += 1
                if get(row, "出貨單號"):
                    shipment_shapes[ship_type][shape(get(row, "出貨單號"))] += 1
            file_stats.append((latest, orders, ship_types))

        file_stats.sort(key=lambda stats: stats[0], reverse=True)  # 最新的檔案在前
        shares = Counter()
        for _, orders, _ in file_stats:
            total = sum(orders.values())
            for platform, count in orders.items():
                shares[platform] += count * 100 / total / len(file_stats)

        def newest_ship_types(platform):
            return next(types[platform] for _, orders, types in file_stats if orders[platform])

        catalog = {
            "platforms": [
                {
                    "name": platform,
                    "share": round(share, 4),  # 平台佔比（%），各檔案平均
                    "shops": sorted(s for s in shops[platform] if s),
                    "ship_types": [[t, n] for t, n in newest_ship_types(platform).most_common() if t],
                    # 範例裡這個平台的建立時間幾乎都是 00:00:00（只有日期），生成時也照做
                    "date_only": date_only[platform][True] >= 0.9 * sum(date_only[platform].values()),
                    "order_no_shapes": [s for s, _ in order_no_shapes[platform].most_common(SHAPE_LIMIT)],
                }
                for platform, share in shares.most_common()
            ],
            "shipment_shapes": {t: [s for s, _ in c.most_common(SHAPE_LIMIT)] for t, c in shipment_shapes.items()},
            "products": [
                [code, name, spec, int(float(prices[(code, name, spec)].most_common(1)[0][0])) if prices[(code, name, spec)] else 0]
                for (code, name, spec), _ in products.most_common(PRODUCT_LIMIT)
            ],
        }
        Path(CATALOG_PATH).write_text(json.dumps(catalog, ensure_ascii=False, indent=1), encoding="utf-8")
        self.stdout.write(self.style.SUCCESS(
            f"已寫入 {CATALOG_PATH}：{len(catalog['platforms'])} 個平台、{len(catalog['products'])} 項商品。"
        ))
