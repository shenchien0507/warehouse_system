"""讀取來源檔（xlsx、csv），依欄位對應版本轉成系統欄位。目前只用來檢查對應，不存訂單。"""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from .fields import FIELDS, FIELDS_BY_KEY

PREVIEW_ROWS = 20  # 測試結果顯示的列數
CHECK_ROWS = 500  # 測試時檢查轉換錯誤的列數，大檔不用整份讀完

DATETIME_FORMATS = [
    "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%Y%m%d",
]


class ImportFileError(Exception):
    pass


def read_rows(uploaded, header_row):
    """回傳 (欄名列, 資料列的 iterator)；欄名前後空白會去掉，整列空白的資料列略過。"""
    name = uploaded.name.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
        except Exception as exc:
            raise ImportFileError("檔案打不開，請確認是 Excel 的 .xlsx 檔。") from exc
        rows = workbook.worksheets[0].iter_rows(values_only=True)
    elif name.endswith(".csv"):
        raw = uploaded.read()
        for encoding in ("utf-8-sig", "cp950"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("CSV 的文字編碼看不懂，請另存成 UTF-8。")
        rows = csv.reader(io.StringIO(text))
    else:
        raise ImportFileError("只支援 .xlsx 或 .csv 檔。")

    for _ in range(header_row - 1):
        next(rows, None)
    headers = next(rows, None)
    if headers is None:
        raise ImportFileError(f"第 {header_row} 列沒有資料，請確認範本的「欄名在第幾列」。")
    headers = ["" if h is None else str(h).strip() for h in headers]
    data = (list(r) for r in rows if any(v not in (None, "") for v in r))
    return headers, data


def convert(kind, value):
    """把一格轉成系統欄位的型態；回傳 (值, 錯誤訊息)。空白回傳 None。"""
    if isinstance(value, str):
        value = value.strip()
    if value in (None, ""):
        return None, ""
    if kind == "int":
        if isinstance(value, bool):
            return None, "不是整數"
        if isinstance(value, (int, float)):
            return (int(value), "") if float(value).is_integer() else (None, "不是整數")
        try:
            return int(str(value).replace(",", "")), ""
        except ValueError:
            return None, "不是整數"
    if kind == "datetime":
        if isinstance(value, datetime):
            return value, ""
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day), ""
        for fmt in DATETIME_FORMATS:
            try:
                return datetime.strptime(str(value), fmt), ""
            except ValueError:
                continue
        return None, "看不懂的日期"
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel 的數字欄（例如交易序號）不要變成 10035943.0
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d %H:%M:%S"), ""
    return str(value), ""


@dataclass
class ColumnCheck:
    field: object  # fields.OrderField
    source_column: str = ""
    index: int = None  # 在來源檔的第幾欄，找不到是 None

    @property
    def found(self):
        return self.index is not None


@dataclass
class InspectResult:
    headers: list
    checks: list
    duplicate_columns: list
    unmapped_columns: list
    preview: list = field(default_factory=list)  # 每列 [(值, 錯誤)]，順序同 mapped_checks
    error_counts: dict = field(default_factory=dict)  # 系統欄位 key → 錯誤列數
    error_samples: list = field(default_factory=list)  # [(第幾列, 欄位名稱, 原始值, 錯誤)]
    blank_counts: dict = field(default_factory=dict)  # 系統欄位 key → 空白列數，只供參考
    checked_rows: int = 0
    more_rows: bool = False

    @property
    def missing_must_map(self):
        return [c for c in self.checks if c.field.must_map and not c.found]

    @property
    def missing_optional(self):
        return [c for c in self.checks if c.source_column and not c.field.must_map and not c.found]

    @property
    def mapped_checks(self):
        return [c for c in self.checks if c.found]

    @property
    def error_summary(self):
        return [(FIELDS_BY_KEY[key].label, count) for key, count in self.error_counts.items()]

    @property
    def blank_summary(self):
        return [(FIELDS_BY_KEY[key].label, count) for key, count in self.blank_counts.items()]

    @property
    def ok(self):
        return not self.missing_must_map and not self.error_counts


def read_headers(uploaded, header_row):
    """只讀欄名列，回傳不重複、非空白的欄名（照檔案順序）。"""
    headers, _ = read_rows(uploaded, header_row)
    return list(dict.fromkeys(h for h in headers if h))


def inspect(version, uploaded):
    """用這個版本的對應讀檔：找欄位、轉換前 CHECK_ROWS 列，回傳 InspectResult。"""
    headers, data = read_rows(uploaded, version.header_row)
    mappings = version.mapping_dict()
    positions = {}
    duplicates = []
    for i, header in enumerate(headers):
        if not header:
            continue
        if header in positions:
            duplicates.append(header)
        else:
            positions[header] = i

    checks = []
    for f in FIELDS:
        source = mappings.get(f.key, "")
        checks.append(ColumnCheck(f, source, positions.get(source) if source else None))

    used = {c.source_column for c in checks if c.found}
    result = InspectResult(
        headers=headers,
        checks=checks,
        duplicate_columns=sorted(set(duplicates)),
        unmapped_columns=[h for h in positions if h not in used],
    )
    mapped = result.mapped_checks
    for row_no, row in enumerate(data, start=version.header_row + 1):
        if result.checked_rows >= CHECK_ROWS:
            result.more_rows = True
            break
        result.checked_rows += 1
        cells = []
        for c in mapped:
            raw = row[c.index] if c.index < len(row) else None
            value, error = convert(c.field.kind, raw)
            if not error and value is None:
                result.blank_counts[c.field.key] = result.blank_counts.get(c.field.key, 0) + 1
            if error:
                result.error_counts[c.field.key] = result.error_counts.get(c.field.key, 0) + 1
                if len(result.error_samples) < 20:
                    result.error_samples.append((row_no, c.field.label, "" if raw is None else raw, error))
            cells.append((value, error))
        if len(result.preview) < PREVIEW_ROWS:
            result.preview.append(cells)
    return result
